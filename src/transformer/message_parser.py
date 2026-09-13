"""
Parse a bank's internal message into a flat field inventory
that the FieldMapper can reason over.

Internal message definitions arrive in three shapes in practice: proprietary
XML, JSON from an internal API, and a field specification held as a spreadsheet
or CSV. All three collapse into the same `InternalField` inventory so the rest
of the Transformation Advisor is format agnostic.
"""
from __future__ import annotations

import csv
import io
import json
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, List, Optional

from lxml import etree

# Spreadsheet / CSV column headers recognised as the field path, the sample
# value and the description, in preference order.
_PATH_HEADERS = ("path", "xpath", "field", "field name", "fieldname",
                 "element", "tag", "name", "attribute")
_VALUE_HEADERS = ("value", "sample", "sample value", "example",
                  "example value", "default")
_DESC_HEADERS = ("description", "notes", "comment", "definition", "meaning")


@dataclass
class InternalField:
    name: str           # local element / attribute name
    xpath: str          # XPath relative to document root
    value: str          # text content (stripped)
    parent: str         # immediate parent element name
    is_attribute: bool = False
    sample: str = ""    # first 80 chars of value (for display)
    description: str = ""   # specification description, where supplied

    def __post_init__(self):
        self.sample = (self.value or "")[:80]


def _walk(element: etree._Element, xpath_parts: List[str], results: List[InternalField]) -> None:
    tag = etree.QName(element.tag).localname
    current_path = "/".join(xpath_parts + [tag])

    text = (element.text or "").strip()
    if text:
        parent = xpath_parts[-1] if xpath_parts else ""
        results.append(InternalField(
            name=tag,
            xpath=current_path,
            value=text,
            parent=parent,
        ))

    for attr_name, attr_val in element.attrib.items():
        local_attr = etree.QName(attr_name).localname
        results.append(InternalField(
            name=local_attr,
            xpath=f"{current_path}/@{local_attr}",
            value=attr_val,
            parent=tag,
            is_attribute=True,
        ))

    for child in element:
        _walk(child, xpath_parts + [tag], results)


def parse_xml_fields(xml_content: str) -> List[InternalField]:
    """
    Parse internal bank XML and return a flat list of InternalField records.
    Skips structural wrapper elements that contain no text of their own.
    """
    try:
        root = etree.fromstring(xml_content.encode("utf-8"))
    except etree.XMLSyntaxError as exc:
        raise ValueError(f"Cannot parse XML: {exc}") from exc

    results: List[InternalField] = []
    _walk(root, [], results)
    return results


# ── JSON ───────────────────────────────────────────────────────────────────


def _walk_json(node: Any, path: List[str], results: List[InternalField]) -> None:
    if isinstance(node, dict):
        for key, value in node.items():
            _walk_json(value, path + [str(key)], results)
        return
    if isinstance(node, list):
        # Repeating structures are inventoried once per occurrence so a mapper
        # sees the cardinality, with the index kept in the path.
        for index, value in enumerate(node, 1):
            _walk_json(value, path + [f"[{index}]"], results)
        return

    if node is None or node == "":
        return
    value = "true" if node is True else "false" if node is False else str(node)
    segments = [p for p in path if not p.startswith("[")]
    results.append(InternalField(
        name=segments[-1] if segments else "",
        xpath="/".join(path),
        value=value,
        parent=segments[-2] if len(segments) > 1 else "",
    ))


def parse_json_fields(json_content: str) -> List[InternalField]:
    """Parse an internal JSON message into the same field inventory as XML."""
    try:
        data = json.loads(json_content)
    except json.JSONDecodeError as exc:
        raise ValueError(f"Cannot parse JSON: {exc}") from exc

    results: List[InternalField] = []
    _walk_json(data, [], results)
    if not results:
        raise ValueError("No fields with values found in the JSON message.")
    return results


# ── Field specifications (CSV / XLSX) ──────────────────────────────────────


def _pick(headers: list[str], candidates: tuple[str, ...]) -> Optional[int]:
    normalised = [(h or "").strip().lower() for h in headers]
    for candidate in candidates:
        if candidate in normalised:
            return normalised.index(candidate)
    for index, header in enumerate(normalised):
        if any(candidate in header for candidate in candidates):
            return index
    return None


def _rows_to_fields(rows: list[list[str]]) -> List[InternalField]:
    rows = [r for r in rows if any((c or "").strip() for c in r)]
    if len(rows) < 2:
        raise ValueError(
            "Field specification needs a header row and at least one field row."
        )
    headers, body = rows[0], rows[1:]
    path_col = _pick(headers, _PATH_HEADERS)
    if path_col is None:
        raise ValueError(
            "No field-name column found. Expected a header such as "
            "'Field', 'Path' or 'Element'."
        )
    value_col = _pick(headers, _VALUE_HEADERS)
    desc_col = _pick(headers, _DESC_HEADERS)

    results: List[InternalField] = []
    for row in body:
        def cell(index: Optional[int]) -> str:
            if index is None or index >= len(row):
                return ""
            return str(row[index] or "").strip()

        path = cell(path_col)
        if not path:
            continue
        path = path.replace("\\", "/").strip("/")
        segments = [s for s in path.split("/") if s]
        name = segments[-1] if segments else path
        results.append(InternalField(
            name=name.lstrip("@"),
            xpath=path,
            value=cell(value_col),
            parent=segments[-2] if len(segments) > 1 else "",
            is_attribute=name.startswith("@"),
            description=cell(desc_col),
        ))
    if not results:
        raise ValueError("No field rows found in the specification.")
    return results


def parse_csv_fields(csv_content: str) -> List[InternalField]:
    """Parse a CSV field specification (one row per internal field)."""
    try:
        dialect = csv.Sniffer().sniff(csv_content[:4096], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    rows = [list(r) for r in csv.reader(io.StringIO(csv_content), dialect)]
    return _rows_to_fields(rows)


def parse_xlsx_fields(xlsx_bytes: bytes) -> List[InternalField]:
    """Parse an XLSX field specification, reading the first worksheet."""
    try:
        from openpyxl import load_workbook
    except ImportError as exc:   # pragma: no cover - openpyxl is a dependency
        raise ValueError("openpyxl is required to read .xlsx specifications") from exc

    try:
        workbook = load_workbook(io.BytesIO(xlsx_bytes), read_only=True, data_only=True)
    except Exception as exc:
        raise ValueError(f"Cannot read spreadsheet: {exc}") from exc

    sheet = workbook[workbook.sheetnames[0]]
    rows = [
        ["" if cell is None else str(cell) for cell in row]
        for row in sheet.iter_rows(values_only=True)
    ]
    workbook.close()
    return _rows_to_fields(rows)


def parse_fields(content: bytes, filename: str = "") -> List[InternalField]:
    """
    Parse an internal message or field specification by file extension,
    falling back to sniffing the content when the name gives nothing away.
    """
    suffix = Path(filename).suffix.lower()
    if suffix == ".xlsx":
        return parse_xlsx_fields(content)

    text = content.decode("utf-8-sig", errors="replace")
    if suffix == ".json":
        return parse_json_fields(text)
    if suffix in (".csv", ".tsv", ".txt"):
        return parse_csv_fields(text)
    if suffix == ".xml":
        return parse_xml_fields(text)

    stripped = text.lstrip()
    if stripped.startswith("<"):
        return parse_xml_fields(text)
    if stripped.startswith(("{", "[")):
        return parse_json_fields(text)
    return parse_csv_fields(text)


def fields_to_summary(fields: List[InternalField]) -> str:
    """Return a compact human-readable inventory for display / agent context."""
    lines = [f"{'XPath':<60} {'Value'}", "-" * 100]
    for f in fields:
        lines.append(f"{f.xpath:<60} {f.sample}")
    lines.append(f"\nTotal fields extracted: {len(fields)}")
    return "\n".join(lines)
