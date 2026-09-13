"""
Business-readable impact assessment for a schema comparison.

The JSON and text reports the comparator already produces are engineering
artefacts. A BA or PO reviewing a schema change needs the same findings framed
as business impact: what the score means, which business flows are affected by
each breaking change, and what has to be decided or built before go-live.

Both renderers take the same `ComparisonResult`, so Markdown and Word never
drift apart.
"""
from __future__ import annotations

import io
from datetime import date
from typing import TYPE_CHECKING, Optional

from src.comparator.diff_classifier import ChangeType
from src.utils.helpers import message_family_from_type

if TYPE_CHECKING:  # avoids a circular import at runtime
    from src.comparator.xml_comparator import ComparisonResult, DiffEntry


# ── Business interpretation of the elements the classifier cares about ────────

_BUSINESS_FLOWS: dict[str, str] = {
    "InstdAmt": "Payment instruction — instructed amount",
    "IntrBkSttlmAmt": "Interbank settlement",
    "Amt": "Payment amount handling",
    "Ccy": "Currency handling and FX",
    "IBAN": "Beneficiary / debtor account routing",
    "BIC": "Bank identification and routing",
    "BICFI": "Bank identification and routing",
    "PmtMtd": "Payment method selection",
    "SttlmMtd": "Settlement method",
    "SvcLvl": "Service level / scheme selection (e.g. SEPA)",
    "PmtTpInf": "Payment type and scheme routing",
    "ReqdExctnDt": "Execution date and value dating",
    "IntrBkSttlmDt": "Settlement date and value dating",
    "MndtRltdInf": "Direct debit mandate handling",
    "RmtInf": "Remittance information and reconciliation",
    "Purp": "Purpose codes and regulatory classification",
    "RgltryRptg": "Regulatory reporting",
    "CtrlSum": "Batch control totals and reconciliation",
    "NbOfTxs": "Batch control totals and reconciliation",
    "Cd": "Coded value / scheme code list",
    # Trade finance — guarantees, standby LCs and trade services baselines
    "UdrtkgAmt": "Guarantee amount and bank exposure",
    "LclUdrtkgAmt": "Local undertaking amount and bank exposure",
    "PlusTlrnce": "Guarantee amount tolerance",
    "XpryDtls": "Expiry and claim window",
    "XpryTerms": "Expiry and claim window",
    "AutoXtnsn": "Automatic extension (evergreen) terms",
    "AutomtcAmtVartn": "Automatic amount variation schedule",
    "PdgQty": "Reported quantities and units of measure",
    "GovncRulesAndLaw": "Governing rules, applicable law and jurisdiction",
    "RuleId": "Governing rules (URDG / ISP98 / UCP)",
    "AplblLaw": "Applicable law and jurisdiction",
    "UdrtkgTermsAndConds": "Undertaking wording and payment obligation",
    "UdrtkgWrdg": "Undertaking wording and payment obligation",
    "Applcnt": "Applicant identification and credit exposure",
    "Bnfcry": "Beneficiary identification and entitlement to claim",
    "AdvsgPty": "Advising bank routing",
    "IssncTp": "Issuance type (issue, amend, counter-guarantee request)",
    "ConfInd": "Confirmation of the undertaking",
    "MltplDmndInd": "Demand handling — multiple demands",
    "PrtlDmndInd": "Demand handling — partial demands",
    "TrfInd": "Transferability of the undertaking",
    "PresntnDtls": "Presentation of documents and demands",
    "UndrlygTx": "Underlying trade transaction and contract cover",
    "AnyBIC": "Party identification and routing",
    "UnitOfMeasr": "Reported quantities and units of measure",
    "OrdrdQty": "Reported quantities and units of measure",
    "AccptdQty": "Reported quantities and units of measure",
    "OutsdngQty": "Reported quantities and units of measure",
    "TxSts": "Trade services baseline status",
}

# Leaf names reused across domains: only meaningful once no parent branch matched
_AMBIGUOUS_LEAVES = {"Amt", "Ccy", "Cd"}

_ACTIONS: dict[ChangeType, str] = {
    ChangeType.ADDED: "Source the new value and populate it before go-live — "
                      "messages without it will be rejected.",
    ChangeType.REMOVED: "Confirm no downstream process depends on this element "
                        "and remove it from the outbound mapping.",
    ChangeType.MODIFIED: "Review the value rules and update transformation and "
                         "validation logic to match.",
    ChangeType.ATTRIBUTE_CHANGED: "Update the attribute in the transformation "
                                  "and re-run validation for affected flows.",
    ChangeType.REORDERED: "Verify the target sequence — strict schema "
                          "validators reject out-of-order elements.",
}

_SEVERITY_MEANING: dict[str, str] = {
    "BREAKING": "Must be fixed before production — messages will fail or "
                "settle incorrectly.",
    "WARNING": "Investigate — impact depends on the receiving implementation.",
    "INFO": "Informational — no action expected.",
    "BENIGN": "Safe to ignore (identifiers, timestamps, correlation refs).",
}


def business_flow(entry: "DiffEntry") -> str:
    """
    Business flow affected by a diff, derived from its path.

    Reusable leaf names (`Amt`, `Cd`, `Dt`, ...) mean different things in a
    payment and in a guarantee, so a named parent branch wins over the leaf.
    """
    segments = [s.split("[")[0].lstrip("@") for s in entry.xpath.split("/") if s]
    element = entry.element.lstrip("@")
    if element and (not segments or segments[-1] != element):
        segments.append(element)

    for name in reversed(segments):
        if name in _BUSINESS_FLOWS and name not in _AMBIGUOUS_LEAVES:
            return _BUSINESS_FLOWS[name]
    for name in reversed(segments):
        if name in _BUSINESS_FLOWS:
            return _BUSINESS_FLOWS[name]
    return "General message structure"


def recommended_action(entry: "DiffEntry") -> str:
    """What a delivery team has to do about a diff."""
    return _ACTIONS.get(entry.change_type, "Review the change with the "
                                           "ISO 20022 standards owner.")


def risk_rating(score: float, breaking: int) -> str:
    """Headline risk rating for the executive summary."""
    if breaking == 0:
        return "LOW"
    if score >= 60:
        return "HIGH"
    if score >= 25:
        return "MEDIUM"
    return "LOW-MEDIUM"


def _message_line(result: "ComparisonResult") -> str:
    """Message type(s) compared, with their business domain."""
    a, b = result.message_type_a, result.message_type_b
    types = f"{a} → {b}" if a and b and a != b else (a or b or "unknown")
    family = message_family_from_type(a or b or "")
    return f"{types} ({family})" if family else types


def _executive_summary(result: "ComparisonResult") -> str:
    breaking = len(result.breaking)
    rating = risk_rating(result.breaking_score, breaking)
    basis = ("classified against the message schema"
             if result.schema_aware else
             "classified on ISO 20022 naming rules only (no schema loaded)")
    if breaking:
        headline = (
            f"{breaking} breaking change(s) require remediation before "
            f"production. Overall risk is {rating} with a breaking-change "
            f"score of {result.breaking_score}/100."
        )
    else:
        headline = (
            f"No breaking changes were found. Overall risk is {rating} with a "
            f"breaking-change score of {result.breaking_score}/100."
        )
    return (
        f"{headline} {len(result.diffs)} difference(s) were detected in total "
        f"({len(result.warnings)} warning(s)), {basis}."
    )


# ── Markdown ──────────────────────────────────────────────────────────────────

def markdown_impact_report(
    result: "ComparisonResult", title: str = "Schema Change Impact Assessment"
) -> str:
    lines = [
        f"# {title}",
        "",
        f"**Baseline:** {result.file_a}  ",
        f"**Compared with:** {result.file_b}  ",
        f"**Message:** {_message_line(result)}  ",
        f"**Date:** {date.today().isoformat()}",
        "",
        "## Executive summary",
        "",
        _executive_summary(result),
        "",
        "| Measure | Value |",
        "| --- | --- |",
        f"| Breaking-change score | {result.breaking_score}/100 |",
        f"| Risk rating | {risk_rating(result.breaking_score, len(result.breaking))} |",
        f"| Breaking changes | {len(result.breaking)} |",
        f"| Warnings | {len(result.warnings)} |",
        f"| Total differences | {len(result.diffs)} |",
        f"| Schema-aware classification | {'yes' if result.schema_aware else 'no'} |",
        "",
        "## How the score is calculated",
        "",
    ]
    if result.score_breakdown:
        lines += [
            "| Severity | Changes | Weight each | Points | Score contribution |",
            "| --- | --- | --- | --- | --- |",
        ]
        for row in result.score_breakdown:
            lines.append(
                f"| {row['severity']} | {row['count']} | {row['weight']} | "
                f"{row['points']} | {row['score_contribution']} |"
            )
        lines += [
            "",
            "Score = points awarded ÷ points if every change were BREAKING.",
            "",
        ]
    else:
        lines += ["No differences were detected, so the score is 0.", ""]

    lines += ["## Breaking changes and affected business flows", ""]
    if result.breaking:
        lines += [
            "| Business flow | Element | Change | From | To | Recommended action |",
            "| --- | --- | --- | --- | --- | --- |",
        ]
        for entry in result.breaking:
            lines.append(
                f"| {business_flow(entry)} | `{entry.xpath}` | "
                f"{entry.change_type.value} | {entry.old_value or '—'} | "
                f"{entry.new_value or '—'} | {recommended_action(entry)} |"
            )
    else:
        lines.append("None — no change was classified as BREAKING.")
    lines.append("")

    if result.warnings:
        lines += [
            "## Warnings to investigate",
            "",
            "| Business flow | Element | Change | Recommended action |",
            "| --- | --- | --- | --- |",
        ]
        for entry in result.warnings:
            lines.append(
                f"| {business_flow(entry)} | `{entry.xpath}` | "
                f"{entry.change_type.value} | {recommended_action(entry)} |"
            )
        lines.append("")

    lines += ["## Severity definitions", "", "| Severity | Meaning |", "| --- | --- |"]
    for severity, meaning in _SEVERITY_MEANING.items():
        lines.append(f"| {severity} | {meaning} |")
    lines.append("")

    if result.validation_errors_a or result.validation_errors_b:
        lines += ["## Schema validation findings", ""]
        for label, errors in (
            (result.file_a, result.validation_errors_a),
            (result.file_b, result.validation_errors_b),
        ):
            for err in errors:
                lines.append(f"- **{label}**: {err}")
        lines.append("")

    return "\n".join(lines)


# ── Word ──────────────────────────────────────────────────────────────────────

def word_impact_report(
    result: "ComparisonResult",
    title: str = "Schema Change Impact Assessment",
    prepared_for: Optional[str] = None,
) -> bytes:
    """Render the same assessment as a .docx, returned as bytes."""
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt, RGBColor

    navy = RGBColor(0x1F, 0x4E, 0x79)
    doc = Document()

    heading = doc.add_paragraph()
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = heading.add_run(title)
    run.bold = True
    run.font.size = Pt(20)
    run.font.color.rgb = navy

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta_run = meta.add_run(
        f"{result.file_a}  vs  {result.file_b}\n"
        f"{_message_line(result)}\n{date.today().isoformat()}"
        + (f"\nPrepared for: {prepared_for}" if prepared_for else "")
    )
    meta_run.font.size = Pt(10)

    doc.add_heading("Executive summary", level=1)
    doc.add_paragraph(_executive_summary(result))

    summary_rows = [
        ("Breaking-change score", f"{result.breaking_score}/100"),
        ("Risk rating", risk_rating(result.breaking_score, len(result.breaking))),
        ("Breaking changes", str(len(result.breaking))),
        ("Warnings", str(len(result.warnings))),
        ("Total differences", str(len(result.diffs))),
        ("Schema-aware classification", "yes" if result.schema_aware else "no"),
    ]
    table = doc.add_table(rows=0, cols=2)
    table.style = "Light Grid Accent 1"
    for label, value in summary_rows:
        cells = table.add_row().cells
        cells[0].text = label
        cells[1].text = value

    doc.add_heading("How the score is calculated", level=1)
    if result.score_breakdown:
        score_table = doc.add_table(rows=1, cols=5)
        score_table.style = "Light Grid Accent 1"
        headers = ["Severity", "Changes", "Weight each", "Points", "Contribution"]
        for cell, header in zip(score_table.rows[0].cells, headers):
            cell.text = header
        for row in result.score_breakdown:
            cells = score_table.add_row().cells
            cells[0].text = row["severity"]
            cells[1].text = str(row["count"])
            cells[2].text = str(row["weight"])
            cells[3].text = str(row["points"])
            cells[4].text = str(row["score_contribution"])
        doc.add_paragraph(
            "Score = points awarded ÷ points if every change were BREAKING."
        )
    else:
        doc.add_paragraph("No differences were detected, so the score is 0.")

    doc.add_heading("Breaking changes and affected business flows", level=1)
    if result.breaking:
        breaking_table = doc.add_table(rows=1, cols=5)
        breaking_table.style = "Light Grid Accent 1"
        headers = ["Business flow", "Element", "Change", "From → To", "Action"]
        for cell, header in zip(breaking_table.rows[0].cells, headers):
            cell.text = header
        for entry in result.breaking:
            cells = breaking_table.add_row().cells
            cells[0].text = business_flow(entry)
            cells[1].text = entry.xpath
            cells[2].text = entry.change_type.value
            cells[3].text = f"{entry.old_value or '—'} → {entry.new_value or '—'}"
            cells[4].text = recommended_action(entry)
    else:
        doc.add_paragraph("None — no change was classified as BREAKING.")

    if result.warnings:
        doc.add_heading("Warnings to investigate", level=1)
        warn_table = doc.add_table(rows=1, cols=4)
        warn_table.style = "Light Grid Accent 1"
        headers = ["Business flow", "Element", "Change", "Action"]
        for cell, header in zip(warn_table.rows[0].cells, headers):
            cell.text = header
        for entry in result.warnings:
            cells = warn_table.add_row().cells
            cells[0].text = business_flow(entry)
            cells[1].text = entry.xpath
            cells[2].text = entry.change_type.value
            cells[3].text = recommended_action(entry)

    doc.add_heading("Severity definitions", level=1)
    for severity, meaning in _SEVERITY_MEANING.items():
        doc.add_paragraph(f"{severity} — {meaning}", style="List Bullet")

    if result.validation_errors_a or result.validation_errors_b:
        doc.add_heading("Schema validation findings", level=1)
        for label, errors in (
            (result.file_a, result.validation_errors_a),
            (result.file_b, result.validation_errors_b),
        ):
            for err in errors:
                doc.add_paragraph(f"{label}: {err}", style="List Bullet")

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


__all__ = [
    "business_flow",
    "markdown_impact_report",
    "recommended_action",
    "risk_rating",
    "word_impact_report",
]
