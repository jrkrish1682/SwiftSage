"""
Phase 2 — Transformation Advisor breadth and trust.

Covers the schema-derived target context, the pacs.008/camt.053 mandatory
registries, multi-format ingestion, ISO XPath validation with confidence
gating, and the requirements-document upgrades.
"""
import io
from pathlib import Path

import pytest

from src.transformer import gap_analyzer, mapping_validator
from src.transformer.field_mapper import (
    MAX_MAPPED_FIELDS,
    MappedField,
    select_fields,
    target_context,
)
from src.transformer.message_parser import (
    parse_csv_fields,
    parse_fields,
    parse_json_fields,
    parse_xlsx_fields,
)
from src.transformer.requirements_generator import (
    Provenance,
    _assumptions,
    generate_requirements_doc,
)
from src.transformer.source_classifier import detect_family, mismatch_warning
from src.transformer.target_schema import (
    RENAMED,
    RESOLVED,
    UNCHECKED,
    UNRESOLVED,
    TargetSchema,
)

SAMPLES = Path(__file__).resolve().parents[1] / "data" / "samples" / "internal"


def _mapping(iso_xpath: str, confidence: str = "HIGH",
             mapping_type: str = "DIRECT", source_value: str = "") -> MappedField:
    return MappedField(
        source_field="Field",
        source_xpath="Header/Field",
        source_value=source_value,
        iso20022_xpath=iso_xpath,
        iso20022_element=iso_xpath.rsplit("/", 1)[-1],
        mapping_type=mapping_type,
        confidence=confidence,
        business_rule="",
        notes="",
    )


# ── Target schema outline ──────────────────────────────────────────────────


@pytest.mark.parametrize("message_type", [
    "pain.001.001.09", "pacs.008.001.10", "camt.053.001.10", "tsrv.001.001.01",
])
def test_every_demo_target_has_a_schema_outline(message_type):
    schema = TargetSchema.for_message_type(message_type)
    assert schema is not None, f"no vendored XSD outline for {message_type}"
    assert len(schema) > 100
    assert schema.mandatory_leaves()


def test_mapping_context_describes_the_selected_target():
    context = TargetSchema.for_message_type("pacs.008.001.10").mapping_context()
    assert "pacs.008.001.10" in context
    assert "IntrBkSttlmAmt" in context
    # Cardinality must be visible so the mapper knows what is mandatory
    assert "[1..1]" in context


def test_mapping_context_is_target_specific_not_always_pain001():
    pacs_context, _, pacs_domain = target_context("pacs.008.001.10")
    pain_context, _, _ = target_context("pain.001.001.09")
    tsrv_context, _, tsrv_domain = target_context("tsrv.001.001.01")

    assert pacs_context != pain_context
    assert "IntrBkSttlmAmt" in pacs_context
    assert "CstmrCdtTrfInitn" not in pacs_context
    assert pacs_domain == "payment"
    assert tsrv_domain == "trade finance"
    assert "Udrtkg" in tsrv_context


# ── Mandatory-field registries ─────────────────────────────────────────────


def test_pacs008_uses_its_own_expert_register():
    fields, origin = gap_analyzer.mandatory_fields("pacs.008.001.10")
    assert origin == "expert"
    paths = [f[0] for f in fields]
    assert any("IntrBkSttlmAmt" in p for p in paths)
    assert any("SttlmInf/SttlmMtd" in p for p in paths)
    assert not any("CdtTrfTxInf/Amt/InstdAmt" in p for p in paths)


def test_camt053_uses_its_own_expert_register():
    fields, origin = gap_analyzer.mandatory_fields("camt.053.001.10")
    assert origin == "expert"
    paths = [f[0] for f in fields]
    assert any("Bal/Amt" in p for p in paths)
    assert any("Stmt/Acct/Id/IBAN" in p for p in paths)


def test_unknown_version_falls_back_within_its_own_family():
    pacs_fields, _ = gap_analyzer.mandatory_fields("pacs.008.001.99")
    expected, _ = gap_analyzer.mandatory_fields("pacs.008.001.10")
    assert pacs_fields == expected


def test_target_without_expert_table_derives_from_the_schema():
    fields, origin = gap_analyzer.mandatory_fields("tsmt.011.001.04")
    assert origin == "schema"
    assert fields
    assert all(len(entry) == 5 for entry in fields)


def test_gap_entries_record_their_origin():
    mappings = [_mapping("GrpHdr/MsgId")]
    gaps = gap_analyzer.analyze(mappings, "camt.053.001.10")
    assert gaps
    assert {g.origin for g in gaps} == {"expert"}


# ── Ingestion ──────────────────────────────────────────────────────────────


def test_json_message_flattens_into_the_field_inventory():
    fields = parse_json_fields(
        (SAMPLES / "sample_bank_statement.json").read_text(encoding="utf-8")
    )
    paths = {f.xpath: f.value for f in fields}
    assert paths["account/accountNumber"] == "20139847"
    assert paths["balances/closingBalance/amount"] == "2015788.13"
    # Repeating structures keep their index so cardinality is visible
    assert "transactions/[1]/entryRef" in paths
    assert "transactions/[3]/entryRef" in paths
    by_name = [f for f in fields if f.xpath == "account/accountNumber"][0]
    assert by_name.name == "accountNumber"
    assert by_name.parent == "account"


def test_json_scalars_are_stringified_and_empties_dropped():
    fields = parse_json_fields(
        '{"a": {"flag": true, "count": 3, "empty": "", "missing": null}}'
    )
    assert {f.xpath: f.value for f in fields} == {
        "a/flag": "true",
        "a/count": "3",
    }


def test_malformed_json_is_rejected_with_a_clear_error():
    with pytest.raises(ValueError, match="Cannot parse JSON"):
        parse_json_fields("{not json")


def test_csv_field_specification_becomes_fields_with_descriptions():
    fields = parse_csv_fields(
        (SAMPLES / "sample_bank_fi_transfer_spec.csv").read_text(encoding="utf-8")
    )
    by_path = {f.xpath: f for f in fields}
    assert by_path["Transfer/SettlementAmount"].value == "312400.00"
    assert by_path["Transfer/Beneficiary/IBAN"].name == "IBAN"
    assert by_path["Transfer/Beneficiary/IBAN"].parent == "Beneficiary"
    assert "Debtor name" in by_path["Transfer/OrderingCustomer/Name"].description
    # A field with no sample value is still inventoried
    assert by_path["Transfer/UETR"].value == ""


def test_csv_header_synonyms_and_semicolon_delimiter_are_accepted():
    fields = parse_csv_fields(
        "Element;Example Value;Notes\n"
        "Header/Ref;ABC123;Internal reference\n"
    )
    assert fields[0].xpath == "Header/Ref"
    assert fields[0].value == "ABC123"
    assert fields[0].description == "Internal reference"


def test_csv_without_a_field_column_is_rejected():
    with pytest.raises(ValueError, match="field-name column"):
        parse_csv_fields("Foo,Bar\n1,2\n")


def test_xlsx_field_specification_is_read_from_the_first_sheet():
    openpyxl = pytest.importorskip("openpyxl")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Field", "Sample Value", "Description"])
    ws.append(["Header/InternalRef", "FIT-2026-0311", "Internal reference"])
    ws.append(["Transfer/Amount", 312400.0, "Settled amount"])
    buffer = io.BytesIO()
    wb.save(buffer)

    fields = parse_xlsx_fields(buffer.getvalue())
    assert [f.xpath for f in fields] == ["Header/InternalRef", "Transfer/Amount"]
    assert fields[1].value == "312400"


def test_parse_fields_routes_by_extension_and_by_content():
    xml = "<Msg><Ref>A1</Ref></Msg>"
    assert parse_fields(xml.encode(), "internal.xml")[0].value == "A1"
    assert parse_fields(xml.encode(), "no-extension")[0].value == "A1"
    assert parse_fields(b'{"Ref": "A1"}', "internal.json")[0].value == "A1"
    assert parse_fields(b"Field,Value\nRef,A1\n", "spec.csv")[0].value == "A1"


# ── Field budget ───────────────────────────────────────────────────────────


@pytest.mark.parametrize("sample", [
    "sample_bank_payment.xml",
    "sample_bank_fi_transfer.xml",
    "sample_bank_fi_transfer_spec.csv",
    "sample_bank_statement.json",
    "sample_bank_guarantee.xml",
])
def test_no_run_sends_more_than_the_budget_and_nothing_is_lost(sample):
    fields = parse_fields((SAMPLES / sample).read_bytes(), sample)
    selected, deferred = select_fields(fields)
    assert len(selected) == MAX_MAPPED_FIELDS
    assert len(selected) + len(deferred) == len(
        [f for f in fields if not f.is_attribute]
    )
    order = {id(f): i for i, f in enumerate(fields)}
    positions = [order[id(f)] for f in selected]
    assert positions == sorted(positions)   # document order preserved


def test_the_budget_keeps_the_fields_a_mapping_cannot_do_without():
    sample = "sample_bank_payment.xml"
    fields = parse_fields((SAMPLES / sample).read_bytes(), sample)
    kept = {f.xpath.rsplit("/", 1)[-1] for f in select_fields(fields)[0]}
    assert {"Amount", "Currency", "ValueDate", "BIC", "SortCode",
            "AccountNumber", "RemittanceInfo"} <= kept


def test_internal_control_fields_lose_their_place_to_business_fields():
    sample = "sample_bank_payment.xml"
    fields = parse_fields((SAMPLES / sample).read_bytes(), sample)
    deferred = {f.xpath.rsplit("/", 1)[-1] for f in select_fields(fields)[1]}
    assert {"WorkflowId", "ApprovalStatus", "ApprovedBy", "CostCentre",
            "Channel"} <= deferred


def test_repeated_occurrences_do_not_consume_the_budget_twice():
    fields = parse_json_fields(
        '{"header": {"ref": "A1"}, "transactions": ['
        + ", ".join(
            '{"amount": "%d.00", "currency": "GBP"}' % n for n in range(1, 31)
        )
        + "]}"
    )
    selected, deferred = select_fields(fields)
    shapes = [f.xpath.replace("[1]", "[]").replace("[2]", "[]") for f in selected]
    assert len([s for s in shapes if s == "transactions/[]/amount"]) == 1
    assert len(deferred) == len(fields) - len(selected)


def test_a_small_message_is_sent_whole():
    fields = parse_csv_fields("Field,Value\nRef,A1\nAmount,10.00\n")
    selected, deferred = select_fields(fields)
    assert len(selected) == 2 and deferred == []


def test_the_budget_is_recorded_as_an_assumption_and_in_provenance():
    provenance = Provenance(field_count=46, fields_mapped=20)
    first = _assumptions([], [], "pain.001.001.09", provenance)[0]
    assert "20 most business-significant of 46" in first
    assert "Gap analysis covers the mapped subset only" in first
    rows = dict(provenance.rows("pain.001.001.09"))
    assert rows["Source fields mapped in this run"] == "20 of 46"

    whole = Provenance(field_count=12, fields_mapped=12)
    assert "business-significant" not in " ".join(
        _assumptions([], [], "pain.001.001.09", whole)
    )
    assert dict(whole.rows("pain.001.001.09"))[
        "Source fields mapped in this run"
    ] == "all inventoried fields"


# ── Source family detection ────────────────────────────────────────────────


@pytest.mark.parametrize("sample,expected_family,target", [
    ("sample_bank_payment.xml", "pain", "pain.001.001.09"),
    ("sample_bank_fi_transfer.xml", "pacs", "pacs.008.001.10"),
    ("sample_bank_fi_transfer_spec.csv", "pacs", "pacs.008.001.10"),
    ("sample_bank_statement.json", "camt", "camt.053.001.10"),
    ("sample_bank_guarantee.xml", "tsrv", "tsrv.001.001.01"),
])
def test_each_demo_source_is_recognised_and_matches_its_target(
    sample, expected_family, target
):
    fields = parse_fields((SAMPLES / sample).read_bytes(), sample)
    family, evidence = detect_family(fields)
    assert family == expected_family
    assert evidence
    assert mismatch_warning(family, target, evidence) == ""


def test_cross_family_target_is_called_out_explicitly():
    sample = "sample_bank_guarantee.xml"
    family, evidence = detect_family(
        parse_fields((SAMPLES / sample).read_bytes(), sample)
    )
    note = mismatch_warning(family, "camt.053.001.10", evidence)
    assert "trade finance undertaking" in note
    assert "camt.053.001.10" in note
    assert "cash management" in note
    assert "guarantee" in note   # the evidence is quoted back


def test_inconclusive_source_is_not_accused_of_a_mismatch():
    fields = parse_csv_fields("Field,Value\nA/B,1\nC/D,2\n")
    assert detect_family(fields) == (None, "")
    assert mismatch_warning(None, "camt.053.001.10") == ""


def test_trade_families_are_treated_as_compatible():
    assert mismatch_warning("tsrv", "tsmt.011.001.04") == ""


def test_family_mismatch_is_recorded_as_an_assumption_in_the_trd():
    provenance = Provenance(source_family="tsrv")
    text = " ".join(_assumptions([], [], "camt.053.001.10", provenance))
    assert "family mismatch" in text
    assert text.index("family mismatch") < 200   # stated first, not buried


def test_provenance_records_the_detected_family():
    rows = dict(Provenance(source_family="pacs").rows("pacs.008.001.10"))
    assert "interbank" in rows["Source message family"]
    rows = dict(Provenance().rows("pacs.008.001.10"))
    assert rows["Source message family"] == "could not be determined"


# ── Mapping validation and confidence gating ───────────────────────────────


@pytest.mark.parametrize("iso_xpath", [
    "PmtInf/CdtTrfTxInf/Amt/InstdAmt",
    # A path relative to the message root resolves just as well as the full one
    "CstmrCdtTrfInitn/PmtInf/CdtTrfTxInf/Amt/InstdAmt",
    "CdtTrfTxInf/Amt/InstdAmt",
])
def test_exact_path_is_confirmed_and_keeps_its_confidence(iso_xpath):
    mapped = mapping_validator.validate([_mapping(iso_xpath)], "pain.001.001.09")[0]
    assert mapped.validation == RESOLVED
    assert mapped.confidence == "HIGH"
    assert mapped.validation_note == ""


def test_element_at_a_different_level_is_flagged_and_downgraded():
    # InstdAmt exists in pain.001, but not under GrpHdr
    mapped = mapping_validator.validate(
        [_mapping("GrpHdr/InstdAmt")], "pain.001.001.09"
    )[0]
    assert mapped.validation == RENAMED
    assert mapped.confidence == "MEDIUM"
    assert "not at the proposed path" in mapped.validation_note
    assert mapped.resolved_xpath


def test_path_absent_from_the_schema_is_downgraded_to_low():
    mapped = mapping_validator.validate(
        [_mapping("GrpHdr/CostCentreCode")], "pain.001.001.09"
    )[0]
    assert mapped.validation == UNRESOLVED
    assert mapped.confidence == "LOW"
    assert "not declared" in mapped.validation_note
    assert mapped.resolved_xpath == ""


def test_unmapped_fields_are_not_validated():
    mapped = mapping_validator.validate(
        [_mapping("", mapping_type="UNMAPPED")], "pain.001.001.09"
    )[0]
    assert mapped.validation == UNCHECKED
    assert mapped.confidence == "HIGH"


def test_no_vendored_schema_leaves_mappings_unchecked():
    mapped = mapping_validator.validate(
        [_mapping("GrpHdr/MsgId")], "pacs.009.001.08"
    )[0]
    assert mapped.validation == UNCHECKED
    assert "No vendored schema" in mapped.validation_note


def test_value_outside_the_target_code_list_is_warned_about():
    mapped = mapping_validator.validate(
        [_mapping("PmtInf/CdtTrfTxInf/ChrgBr", source_value="SHARED")],
        "pain.001.001.09",
    )[0]
    assert mapped.validation == RESOLVED
    assert "not one of the permitted codes" in mapped.validation_note

    ok = mapping_validator.validate(
        [_mapping("PmtInf/CdtTrfTxInf/ChrgBr", source_value="SHAR")],
        "pain.001.001.09",
    )[0]
    assert ok.validation_note == ""


def test_summarise_counts_each_status():
    mappings = mapping_validator.validate(
        [
            _mapping("PmtInf/CdtTrfTxInf/Amt/InstdAmt"),
            _mapping("GrpHdr/CostCentreCode"),
            _mapping("", mapping_type="UNMAPPED"),
        ],
        "pain.001.001.09",
    )
    counts = mapping_validator.summarise(mappings)
    assert counts[RESOLVED] == 1
    assert counts[UNRESOLVED] == 1
    assert counts[UNCHECKED] == 1


# ── Requirements document ──────────────────────────────────────────────────


def test_assumptions_reflect_what_the_run_actually_did():
    mappings = mapping_validator.validate(
        [
            _mapping("GrpHdr/CostCentreCode"),
            _mapping("PmtInf/DbtrAgt/FinInstnId/BICFI", mapping_type="DERIVED"),
            _mapping("", mapping_type="UNMAPPED"),
        ],
        "pain.001.001.09",
    )
    text = " ".join(_assumptions(mappings, [], "pain.001.001.09"))
    assert "pain.001.001.09" in text
    assert "DERIVED" in text
    assert "could not be found in the schema" in text
    assert "unmapped internal field" in text


def test_schema_derived_register_is_declared_as_an_assumption():
    gaps = gap_analyzer.analyze([_mapping("Foo/Bar")], "tsmt.011.001.04")
    text = " ".join(_assumptions([], gaps, "tsmt.011.001.04"))
    assert "derived from the" in text


def test_provenance_hash_is_stable_and_content_sensitive():
    assert Provenance.hash_input("abc") == Provenance.hash_input(b"abc")
    assert Provenance.hash_input("abc") != Provenance.hash_input("abd")
    assert len(Provenance.hash_input("abc")) == 16


def test_requirements_doc_carries_traceability_and_provenance():
    from docx import Document

    mappings = mapping_validator.validate(
        [
            _mapping("PmtInf/CdtTrfTxInf/Amt/InstdAmt"),
            _mapping("GrpHdr/CostCentreCode"),
        ],
        "pain.001.001.09",
    )
    gaps = gap_analyzer.analyze(mappings, "pain.001.001.09")
    provenance = Provenance(
        model="claude-sonnet-4-6",
        source_name="sample_bank_payment.xml",
        source_format="XML",
        input_hash=Provenance.hash_input("payload"),
        schema_source="vendored XSD, 1061 paths",
        field_count=46,
    )

    doc_bytes = generate_requirements_doc(
        mappings, gaps, "pain.001.001.09", provenance=provenance
    )
    doc = Document(io.BytesIO(doc_bytes))
    text = "\n".join(p.text for p in doc.paragraphs)
    cells = [c.text for t in doc.tables for r in t.rows for c in r.cells]

    assert "Assumptions" in text
    assert "Provenance" in text
    assert any(c.startswith("AS-01") or "AS-01" in c for c in [text])
    assert "TR-01" in cells and "TR-02" in cells
    assert any(c.startswith("GAP-01") for c in cells)
    assert "Schema Check" in cells
    assert "Confirmed in schema" in cells
    assert "Not in schema" in cells
    assert provenance.input_hash in cells
    assert "claude-sonnet-4-6" in cells
    assert "sample_bank_payment.xml" in cells


def test_requirements_doc_still_generates_without_provenance():
    doc_bytes = generate_requirements_doc(
        [_mapping("GrpHdr/MsgId")], [], "tsrv.001.001.01"
    )
    assert doc_bytes.startswith(b"PK")
